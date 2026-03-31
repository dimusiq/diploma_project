"""
Собирает пример .xlsx для импорта техники (первый лист, строка заголовков).

Перед загрузкой в приложение замените значения в столбце «Бренд» на реальные
имена из справочника «Администрирование → Бренды» (или создайте эти бренды).

Запуск из каталога backend:
  uv sync && uv run python scripts/generate_equipment_import_sample_xlsx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

BACKEND_ROOT = Path(__file__).resolve().parent.parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.services.equipment_import import IMPORT_TEMPLATE_HEADERS  # noqa: E402

# backend/scripts → репозиторий nebardak
REPO_ROOT = Path(__file__).resolve().parent.parent.parent
OUT_PATH = REPO_ROOT / "scripts" / "sample-data" / "equipment_import_sample.xlsx"

HEADERS = list(IMPORT_TEMPLATE_HEADERS)

# Демо-бренды — должны существовать в БД или замените на свои
ROWS: list[list[object]] = [
    [
        "Электропогрузчик",
        "DemoBrand",
        "E20-01",
        "",
        "SN-DEMO-1001",
        "Г-201",
        "2024-01-15",
        850,
        "В эксплуатации",
        "Приёмка",
        "",
        "Пример строки для теста импорта",
    ],
    [
        "Ричтрак",
        "DemoBrand",
        "RR 1.6",
        "",
        "SN-DEMO-1002",
        "Г-202",
        "2023-08-20",
        3100,
        "active",
        "Высотное хранение",
        "",
        "",
    ],
    [
        "Автопогрузчик",
        "DemoBrand",
        "FG25",
        "VIN-DEMO-003",
        "SN-DEMO-1003",
        "Г-203",
        "15.03.2022",
        5200,
        "На обслуживании",
        "Отгрузка",
        "",
        "",
    ],
    [
        "Комплектовщик",
        "DemoBrand",
        "OPX 20",
        "",
        "SN-DEMO-1004",
        "Г-204",
        "",
        "",
        "В эксплуатации",
        "Комплектация",
        "",
        "",
    ],
    [
        "Электротележка",
        "DemoBrand",
        "EJE 120",
        "",
        "SN-DEMO-1005",
        "Г-205",
        "2025-02-01",
        42,
        "В эксплуатации",
        "Кросс-док",
        "",
        "",
    ],
]


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Техника"

    header_font = Font(bold=True)
    header_align = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    body_align = Alignment(wrap_text=True, vertical="top")

    for idx, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.font = header_font
        cell.alignment = header_align
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = min(28, max(12, len(str(title)) + 4))

    for r, row in enumerate(ROWS, start=2):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = body_align

    wb.save(OUT_PATH)
    print(f"Записано: {OUT_PATH}")


if __name__ == "__main__":
    main()
