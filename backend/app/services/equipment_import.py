"""Разбор Excel .xlsx и подготовка строк для импорта техники."""

from __future__ import annotations

import json
import re
import uuid
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.utils.datetime import from_excel  # type: ignore[import-untyped]
from sqlmodel import Session, select

from app.models import EQUIPMENT_TYPES, Brand, Equipment, EquipmentCreate

MAX_IMPORT_ROWS = 2000
MAX_IMPORT_BYTES = 8 * 1024 * 1024

# Первая строка шаблона импорта и эталонные подписи столбцов (русские названия)
IMPORT_TEMPLATE_HEADERS: tuple[str, ...] = (
    "Тип техники",
    "Бренд",
    "Модель",
    "VIN",
    "Серийный номер",
    "Гаражный номер",
    "Дата ввода",
    "Моточасы",
    "Состояние",
    "Зона",
    "Прицепное оборудование",
    "Примечания",
)

# Нормализованный заголовок столбца -> внутреннее имя поля
HEADER_TO_FIELD: dict[str, str] = {
    "equipment_type": "equipment_type",
    "тип": "equipment_type",
    "тип техники": "equipment_type",
    "вид техники": "equipment_type",
    "brand": "brand",
    "бренд": "brand",
    "brand_id": "brand",
    "id бренда": "brand",
    "model": "model",
    "модель": "model",
    "vin": "vin",
    "serial_number": "serial_number",
    "серийный номер": "serial_number",
    "серийный": "serial_number",
    "garage_number": "garage_number",
    "гаражный номер": "garage_number",
    "гаражный": "garage_number",
    "commissioned_at": "commissioned_at",
    "дата ввода": "commissioned_at",
    "дата ввода в эксплуатацию": "commissioned_at",
    "engine_hours": "engine_hours",
    "моточасы": "engine_hours",
    "current_status": "current_status",
    "состояние": "current_status",
    "статус": "current_status",
    "zone": "zone",
    "зона": "zone",
    "attachments": "attachments",
    "прицепное оборудование": "attachments",
    "instructions": "instructions",
    "инструкции": "instructions",
    "примечания": "instructions",
}

_TYPE_LABEL_TO_ID: dict[str, str] = {
    "автопогрузчик": "autopogruzchik",
    "электропогрузчик": "elektropogruzchik",
    "комплектовщик": "komplektovshchik",
    "ричтрак": "richtrak",
    "электротележка": "elektrotelezhka",
}

_STATUS_LABEL_TO_ID: dict[str, str] = {
    "в эксплуатации": "active",
    "эксплуатация": "active",
    "active": "active",
    "на обслуживании": "maintenance",
    "обслуживание": "maintenance",
    "maintenance": "maintenance",
    "ремонт": "maintenance",
    "выведена из эксплуатации": "decommissioned",
    "выведена": "decommissioned",
    "decommissioned": "decommissioned",
    "списана": "decommissioned",
}


def _norm_header(val: object) -> str:
    if val is None:
        return ""
    s = str(val).strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _cell_str(val: object) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def _parse_date_val(val: object) -> date | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, (int, float)):
        try:
            dt = from_excel(float(val))
            if isinstance(dt, datetime):
                return dt.date()
            if isinstance(dt, date):
                return dt
        except Exception:
            pass
        return None
    s = str(val).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_engine_hours(val: object) -> int | None:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        n = int(val)
        return n if n >= 0 else None
    s = str(val).strip().replace(",", ".")
    if not s:
        return None
    try:
        n = int(float(s))
        return n if n >= 0 else None
    except ValueError:
        return None


def _parse_equipment_type(raw: str) -> str | None:
    if not raw:
        return None
    s = raw.strip().lower()
    s_compact = re.sub(r"\s+", "", s)
    if s in EQUIPMENT_TYPES:
        return s
    if s_compact in EQUIPMENT_TYPES:
        return s_compact
    mapped = _TYPE_LABEL_TO_ID.get(s) or _TYPE_LABEL_TO_ID.get(s_compact)
    if mapped:
        return mapped
    return None


def _parse_current_status(raw: str) -> str:
    if not raw:
        return "active"
    s = raw.strip().lower()
    s = re.sub(r"\s+", " ", s)
    if s in ("active", "maintenance", "decommissioned"):
        return s
    mapped = _STATUS_LABEL_TO_ID.get(s)
    if mapped:
        return mapped
    return "active"


def _parse_attachments_field(raw: str) -> str | None:
    if not raw.strip():
        return None
    try:
        v = json.loads(raw)
        if isinstance(v, list):
            return json.dumps([str(x).strip() for x in v if str(x).strip()])
    except json.JSONDecodeError:
        pass
    parts = [p.strip() for p in raw.replace(";", "\n").splitlines() if p.strip()]
    return json.dumps(parts) if parts else None


def _load_brand_map(session: Session) -> tuple[dict[str, uuid.UUID], set[uuid.UUID]]:
    brands = list(session.exec(select(Brand)).all())
    by_name: dict[str, uuid.UUID] = {}
    ids: set[uuid.UUID] = set()
    for b in brands:
        ids.add(b.id)
        key = b.name.strip().lower()
        if key and key not in by_name:
            by_name[key] = b.id
    return by_name, ids


def _resolve_brand_id(
    raw: str,
    brand_by_lower: dict[str, uuid.UUID],
    valid_brand_ids: set[uuid.UUID],
) -> uuid.UUID | None:
    if not raw:
        return None
    s = raw.strip()
    u = _try_uuid(s)
    if u is not None:
        return u if u in valid_brand_ids else None
    return brand_by_lower.get(s.lower())


def _try_uuid(s: str) -> uuid.UUID | None:
    try:
        return uuid.UUID(s.strip())
    except (ValueError, AttributeError):
        return None


def _read_xlsx_rows(data: bytes) -> list[list[Any]]:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def read_spreadsheet_rows(data: bytes, filename: str) -> list[list[Any]]:
    name = (filename or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise ValueError("Допустим только формат .xlsx (Excel 2007 и новее).")
    if len(data) < 4 or not data.startswith(b"PK"):
        raise ValueError("Файл не похож на .xlsx (ожидается книга Excel 2007+).")
    return _read_xlsx_rows(data)


def build_header_index(header_row: list[Any]) -> dict[str, int]:
    index: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        key = HEADER_TO_FIELD.get(_norm_header(cell))
        if key is not None:
            index[key] = i
    return index


def row_to_field_dict(
    row: list[Any], header_index: dict[str, int]
) -> dict[str, str]:
    out: dict[str, str] = {}
    for field, col in header_index.items():
        if col < len(row):
            v = row[col]
            if field == "commissioned_at":
                d = _parse_date_val(v)
                out[field] = d.isoformat() if d else ""
            elif field == "engine_hours":
                h = _parse_engine_hours(v)
                out[field] = str(h) if h is not None else ""
            else:
                out[field] = _cell_str(v)
        else:
            out[field] = ""
    return out


def fields_to_create(
    fields: dict[str, str],
    brand_by_lower: dict[str, uuid.UUID],
    valid_brand_ids: set[uuid.UUID],
) -> tuple[EquipmentCreate | None, str | None]:
    model = (fields.get("model") or "").strip()
    brand_raw = (fields.get("brand") or "").strip()
    eq_type_raw = (fields.get("equipment_type") or "").strip()

    if not model and not brand_raw and not eq_type_raw:
        return None, None

    if not eq_type_raw:
        return None, "Не указан тип техники"
    eq_type = _parse_equipment_type(eq_type_raw)
    if not eq_type:
        return None, f"Неизвестный тип техники: {eq_type_raw!r}"

    if not brand_raw:
        return None, "Не указан бренд"
    brand_id = _resolve_brand_id(brand_raw, brand_by_lower, valid_brand_ids)
    if brand_id is None:
        return None, f"Бренд не найден в справочнике: {brand_raw!r}"

    if not model:
        return None, "Не указана модель"

    commissioned = None
    comm_raw = (fields.get("commissioned_at") or "").strip()
    if comm_raw:
        commissioned = _parse_date_val(comm_raw)
        if commissioned is None:
            return None, "Некорректная дата ввода в эксплуатацию"

    engine_hours = _parse_engine_hours(
        fields["engine_hours"] if fields.get("engine_hours") else None
    )
    status = _parse_current_status(fields.get("current_status") or "")

    attachments = _parse_attachments_field(fields.get("attachments") or "")
    instructions = (fields.get("instructions") or "").strip() or None

    vin = (fields.get("vin") or "").strip() or None
    serial_number = (fields.get("serial_number") or "").strip() or None
    garage_number = (fields.get("garage_number") or "").strip() or None
    zone = (fields.get("zone") or "").strip() or None

    try:
        body = EquipmentCreate(
            equipment_type=eq_type,
            vin=vin,
            serial_number=serial_number,
            garage_number=garage_number,
            brand_id=brand_id,
            model=model[:128],
            commissioned_at=commissioned,
            engine_hours=engine_hours,
            current_status=status,
            zone=zone,
            attachments=attachments,
            instructions=instructions[:2048] if instructions else None,
        )
    except Exception as e:
        return None, str(e)
    return body, None


def import_equipment_from_spreadsheet(
    session: Session, data: bytes, filename: str
) -> tuple[int, list[tuple[int, str]]]:
    """
    Создаёт записи техники построчно. Возвращает (число созданных, список (номер строки Excel, текст ошибки)).
    """
    rows = read_spreadsheet_rows(data, filename)
    if not rows:
        return 0, [(1, "Файл пустой")]

    header_index = build_header_index(rows[0])
    required_headers = ("equipment_type", "brand", "model")
    missing = [h for h in required_headers if h not in header_index]
    if missing:
        return 0, [
            (
                1,
                "В первой строке должны быть столбцы: тип техники, бренд, модель "
                f"(не найдены: {', '.join(missing)})",
            )
        ]

    brand_by_lower, valid_brand_ids = _load_brand_map(session)
    if not valid_brand_ids:
        return 0, [(1, "Справочник брендов пуст — добавьте бренды в администрировании")]

    data_rows = rows[1 : 1 + MAX_IMPORT_ROWS]
    created = 0
    errors: list[tuple[int, str]] = []

    for i, row in enumerate(data_rows):
        excel_row = i + 2
        fields = row_to_field_dict(row, header_index)
        body, err = fields_to_create(fields, brand_by_lower, valid_brand_ids)
        if body is None:
            if err is None:
                continue
            errors.append((excel_row, err))
            continue
        if body.equipment_type not in EQUIPMENT_TYPES:
            errors.append((excel_row, "Недопустимый тип техники"))
            continue
        if session.get(Brand, body.brand_id) is None:
            errors.append((excel_row, "Бренд не найден"))
            continue
        eq = Equipment.model_validate(body)
        session.add(eq)
        try:
            session.commit()
            session.refresh(eq)
            created += 1
        except Exception as e:
            session.rollback()
            errors.append((excel_row, f"Ошибка сохранения: {e!s}"))

    return created, errors


def build_equipment_import_template_xlsx() -> bytes:
    """Пустой .xlsx: первый лист, одна строка — названия заполняемых полей."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Техника"
    header_font = Font(bold=True)
    header_align = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    for idx, title in enumerate(IMPORT_TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.font = header_font
        cell.alignment = header_align
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = min(42, max(16, len(title) + 4))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
